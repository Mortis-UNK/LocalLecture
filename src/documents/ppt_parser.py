from pathlib import Path

from src.generation.schemas import Slide


def shape_text(shape):
    if getattr(shape, "has_text_frame", False):
        yield shape.text
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            yield " | ".join(cell.text for cell in row.cells)
    if hasattr(shape, "shapes"):
        for child in shape.shapes:
            yield from shape_text(child)


def parse_pptx(path: str) -> list[Slide]:
    from pptx import Presentation

    return [Slide(id=f"slide:{i}", source_file=Path(path).name, page=i,
                  text="\n".join(t for shape in slide.shapes for t in shape_text(shape) if t.strip()))
            for i, slide in enumerate(Presentation(str(path)).slides, start=1)]
