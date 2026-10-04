import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from src.demo import ExtractiveGenerator
from src.generation.schemas import Settings
from src.pipeline import run


class DocumentOnlyTests(unittest.TestCase):
    def test_pptx_and_pdf_without_audio(self):
        from pptx import Presentation
        import pymupdf

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            ppt = Presentation()
            ppt.slides.add_slide(ppt.slide_layouts[6])  # Empty page must not shift citations.
            slide = ppt.slides.add_slide(ppt.slide_layouts[1])
            slide.shapes.title.text = "Gradient descent"
            slide.placeholders[1].text = "Gradient descent reduces the loss."
            ppt.save(root / "lecture.pptx")
            with pymupdf.open() as pdf:
                pdf.new_page()
                page = pdf.new_page()
                page.insert_text((72, 72), "Gradient descent reduces the loss.")
                pdf.save(root / "lecture.pdf")
            settings = Settings(output_dir=str(root / "notes"), processed_dir=str(root / "processed"))
            for extension in ("pptx", "pdf"):
                with self.subTest(extension=extension), patch("src.pipeline.OllamaGenerator", return_value=ExtractiveGenerator()), \
                        patch("src.pipeline.SentenceEncoder") as encoder, patch("src.pipeline.transcribe") as transcriber:
                    markdown, _, path = run(slides_path=root / f"lecture.{extension}", settings=settings)
                    data = json.loads(Path(path).read_text(encoding="utf-8"))
                    self.assertEqual(data["input_mode"], "slides_only")
                    self.assertIn("Slide 2", markdown)
                    self.assertNotIn("🎧", markdown)
                    self.assertTrue(data["notes"][0]["claims"])
                    self.assertTrue(all(r["type"] == "slide" for r in data["sources"].values()))
                    encoder.assert_not_called()
                    transcriber.assert_not_called()
            self.assertFalse(list((root / "processed").rglob("transcript.json")))

    def test_missing_inputs_have_actionable_message(self):
        with self.assertRaisesRegex(ValueError, "请上传"):
            run()

    def test_empty_deck_fails_before_model_loading(self):
        with patch("src.pipeline.parse_pptx", return_value=[]), patch("src.pipeline.OllamaGenerator") as generator:
            with self.assertRaisesRegex(ValueError, "没有可提取的文字"):
                run(slides_path="empty.pptx")
            generator.assert_not_called()
